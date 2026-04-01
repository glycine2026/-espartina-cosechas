import io
import re
import pathlib
import pandas as pd
from openpyxl import load_workbook

BASE_DIR = pathlib.Path(__file__).parent.parent
TEMPLATE_PATH = BASE_DIR / "assets" / "ImportacionCosecha_template.xlsx"

CAMPOS_CON_CODIGO = {
    "Código tipo de comprobante","Tipo CPE","Numerador CPE","Código socio",
    "Código campaña","Código especie","Tipo de Grano","Código cultivo",
    "Código depósito origen","Código depósito destino","Número comprobante contrato",
    "Código contratista","Código embolsador","Código transportista",
    "Código intermediario flete","Código entregador","Código remitente",
    "Código destinatario/mercado a termino","Código destino","Tipo de Flete",
    "Código CATAC","Código corredor","Código extractor",
    "Código Personal Chofer","Código Camión Personal","Pagador Flete",
}

MAPA = {
    "Turno":"Turno","Fecha":"Fecha",
    "Código socio":"Código socio","Código campaña":"Código campaña",
    "Código especie":"Código especie","Código cultivo":"Código cultivo",
    "Código depósito origen":"Código depósito origen",
    "Código depósito destino":"Código depósito destino",
    "Número comprobante contrato":"Número comprobante contrato",
    "Pagador Flete":"Pagador Flete","CUIT Pagador Flete":"CUIT Pagador Flete",
    "Peso Origen Bruto":"Peso Origen Bruto","Peso Origen Tara":"Peso Origen Tara",
    "Peso Origen Neto":"Peso Origen Neto","% Humedad Origen":"% Humedad Origen",
    "% Humedad Destino":"% Humedad Destino","Peso Destino Bruto":"Peso Destino Bruto",
    "Peso Destino Tara":"Peso Destino Tara","Peso Destino Neto":"Peso Destino Neto",
    "Código embolsador":"Código embolsador","Precio Embolsador por TN":"Precio Embolsador por TN",
    "Chofer (CUIT)":"Chofer (CUIT)","Código Personal Chofer":"Código Personal Chofer",
    "Código transportista":"Código transportista",
    "Código intermediario flete":"Código intermediario flete",
    "Tipo CPE":"Tipo CPE","Numerador CPE":"Numerador CPE",
    "Carta de Porte":"Carta de Porte","CTG":"CTG",
    "Catac":"Código CATAC","Distancia Planta":"Distancia Planta",
    "Tarifa Catac":"Tarifa Catac","Coeficiente":"Coeficiente","Aforado":"Aforado",
    "Código destino":"Código destino","Código extractor":"Código extractor",
    "Precio Extractor por TN":"Precio Extractor por TN",
}

def _extraer_codigo(campo, valor):
    if valor is None or str(valor).strip() in ("","nan","None"):
        return ""
    s = str(valor).strip()
    if campo not in CAMPOS_CON_CODIGO:
        return s
    if campo == "Chofer (CUIT)":
        m = re.search(r'\(([^)]+)\)', s)
        return m.group(1) if m else s
    m = re.match(r'^([^\s\-]+)', s)
    return m.group(1) if m else s

def generar_excel(ctgs_data: list) -> bytes:
    wb = load_workbook(str(TEMPLATE_PATH))
    ws = wb["Cosechas a importar"]
    header = [str(c).strip() if c else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col_idx = {h: i+1 for i, h in enumerate(header) if h}

    for row_num, ctg in enumerate(ctgs_data, start=2):
        todos = {**ctg.get("precargado", {}), **ctg.get("formulario", {})}
        ws.cell(row=row_num, column=col_idx.get("Código tipo de comprobante", 1), value="COSI")
        for campo_form, campo_tpl in MAPA.items():
            valor = todos.get(campo_form) or todos.get(campo_tpl, "")
            limpio = _extraer_codigo(campo_tpl, valor)
            if campo_tpl in col_idx and limpio:
                ws.cell(row=row_num, column=col_idx[campo_tpl], value=limpio)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
